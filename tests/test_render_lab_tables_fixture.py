"""Keep the heading-less wide-column fixture isolated from other features."""

import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest import mock

try:
    from openpyxl import load_workbook
    from PIL import Image  # noqa: F401 - fixtures.py imports Pillow
except ImportError:  # the render-lab container supplies these dependencies
    load_workbook = None


@unittest.skipIf(load_workbook is None, "openpyxl and Pillow are the render lab's")
class TablesFixtureTest(unittest.TestCase):
    def test_wide_column_has_no_print_headings_or_cell_furniture(self):
        path = Path(__file__).resolve().parents[1] / "tools/render-lab/fixtures.py"
        spec = importlib.util.spec_from_file_location("render_fixtures", path)
        fixtures = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(fixtures)
        with tempfile.TemporaryDirectory() as out, mock.patch.object(fixtures, "OUT", out):
            fixtures.tables()
            entry = next(e for e in fixtures.MANIFEST if e["feature"] == "wide-column-no-print-headings")
            ws = load_workbook(Path(out) / entry["file"]).active
            self.assertFalse(ws.print_options.headings)
            self.assertFalse(ws.print_options.gridLines)
            self.assertEqual(ws.column_dimensions["A"].width, 80)
            self.assertEqual((ws.max_row, ws.max_column), (4, 1))
            for row in ws.iter_rows():
                cell = row[0]
                self.assertIsInstance(cell.value, str)
                self.assertFalse(cell.alignment.wrap_text)
                self.assertFalse(cell.has_style)
            # The exception must not change the usual heading-framed fixtures.
            values = load_workbook(Path(out) / "tables/values.xlsx").active
            self.assertTrue(values.print_options.headings)
